from types import SimpleNamespace

from app.services import export_service
from app.utils.markdown import clean_markdown_for_export


def _question(**kw):
    base = dict(
        qtype="mcq", question="2+2?", answer="4",
        options=["3", "4", "5", "6"], explanation="Addition.",
    )
    base.update(kw)
    return SimpleNamespace(**base)


def test_clean_markdown_strips_visuals():
    md = (
        "## H\n"
        "```mermaid\nflowchart TD\nA-->B\n```\n"
        '<svg viewBox="0 0 10 10"><text>x</text></svg>\n'
        "[[DIAGRAM: a detailed figure]]\n"
        "text"
    )
    cleaned = clean_markdown_for_export(md)
    assert "<svg" not in cleaned
    assert "```mermaid" not in cleaned
    assert "[[DIAGRAM" not in cleaned
    assert "view in the app" in cleaned


def test_build_pdf_and_docx(fake_chapter):
    questions = [_question(), _question(qtype="true_false", options=[], answer="True")]
    pdf = export_service.build_pdf(fake_chapter, questions, True, True)
    docx = export_service.build_docx(fake_chapter, questions, True, True)
    assert pdf.startswith(b"%PDF-")
    assert docx.startswith(b"PK\x03\x04")
    assert len(pdf) > 500 and len(docx) > 500


def test_worksheet_docx_separates_questions_and_answer_key(fake_chapter):
    import io

    from docx import Document

    questions = [_question(question="2+2?", answer="4")]
    data = export_service.build_docx(
        fake_chapter, questions, include_notes=False, include_questions=True, worksheet=True
    )
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    # Worksheet + answer key sections exist; the answer only appears in the key.
    assert "Worksheet" in texts
    assert "Answer Key" in texts
    assert any("2+2?" in t for t in texts)
    assert not any(t.startswith("Answer:") for t in texts)  # no inline answers


def test_worksheet_pdf_builds(fake_chapter):
    questions = [_question()]
    pdf = export_service.build_pdf(
        fake_chapter, questions, include_notes=False, include_questions=True, worksheet=True
    )
    assert pdf.startswith(b"%PDF-") and len(pdf) > 500


def test_references_section_from_sources():
    import io

    from docx import Document

    chapter = SimpleNamespace(
        chapter_name="Cells", class_name="10", subject="Bio",
        notes="## Overview\nText.",
        sources='[{"url": "https://example.org/cells", "title": "Cell Biology"}]',
    )
    data = export_service.build_docx(chapter, [], include_notes=True, include_questions=False)
    texts = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "References" in texts
    assert any("https://example.org/cells" in t for t in texts)


def test_no_references_when_no_sources(fake_chapter):
    # fake_chapter has no `sources` attribute -> no References section.
    data = export_service.build_docx(fake_chapter, [], include_notes=True, include_questions=False)
    import io

    from docx import Document

    texts = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert "References" not in texts
