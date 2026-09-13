"""Render a chapter's notes and question bank to PDF and DOCX."""
from __future__ import annotations

import io
import json

from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from ..models.enums import label_for
from ..utils.markdown import BOLD_RE, clean_markdown_for_export, strip_md

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _inline_html(text: str) -> str:
    """Escape XML then re-apply **bold** as <b> for reportlab paragraphs."""
    out = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return BOLD_RE.sub(r"<b>\1</b>", out)


def _options(q) -> list[str]:
    if isinstance(q.options, list):
        return q.options
    if q.options:
        try:
            return json.loads(q.options)
        except json.JSONDecodeError:
            return []
    return []


def _sources(chapter) -> list[dict]:
    """Parse the chapter's JSON-encoded [{url, title}] web sources."""
    raw = getattr(chapter, "sources", "") or ""
    if isinstance(raw, list):
        data = raw
    elif raw:
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            return []
    else:
        return []
    return [s for s in data if isinstance(s, dict) and s.get("url")]


def _group_by_type(questions) -> dict[str, list]:
    grouped: dict[str, list] = {}
    for q in questions:
        grouped.setdefault(q.qtype, []).append(q)
    return grouped


# ------------------------------- PDF -------------------------------- #

def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("H1x", parent=styles["Heading1"], fontSize=18, spaceAfter=10))
    styles.add(
        ParagraphStyle("H2x", parent=styles["Heading2"], fontSize=14, spaceBefore=8, spaceAfter=6)
    )
    styles.add(
        ParagraphStyle("H3x", parent=styles["Heading3"], fontSize=12, spaceBefore=6, spaceAfter=4)
    )
    styles.add(
        ParagraphStyle(
            "Bodyx", parent=styles["BodyText"], fontSize=10.5, leading=15, alignment=TA_LEFT
        )
    )
    styles.add(
        ParagraphStyle("Metax", parent=styles["BodyText"], fontSize=10, textColor="#555555")
    )
    return styles


def _markdown_to_flowables(md: str, styles) -> list:
    flow: list = []
    bullets: list = []

    def flush():
        nonlocal bullets
        if bullets:
            flow.append(
                ListFlowable(
                    [ListItem(Paragraph(_inline_html(b), styles["Bodyx"])) for b in bullets],
                    bulletType="bullet",
                    leftIndent=14,
                )
            )
            bullets = []

    for raw in md.splitlines():
        stripped = raw.strip()
        if not stripped:
            flush()
            flow.append(Spacer(1, 5))
        elif stripped.startswith("### "):
            flush()
            flow.append(Paragraph(_inline_html(stripped[4:]), styles["H3x"]))
        elif stripped.startswith("## "):
            flush()
            flow.append(Paragraph(_inline_html(stripped[3:]), styles["H2x"]))
        elif stripped.startswith("# "):
            flush()
            flow.append(Paragraph(_inline_html(stripped[2:]), styles["H2x"]))
        elif stripped.startswith(("- ", "* ", "+ ")):
            bullets.append(stripped[2:])
        elif stripped.startswith("> "):
            flush()
            flow.append(Paragraph(f"<i>{_inline_html(stripped[2:])}</i>", styles["Bodyx"]))
        else:
            flush()
            flow.append(Paragraph(_inline_html(stripped), styles["Bodyx"]))
    flush()
    return flow


def _pdf_worksheet(flow: list, questions, styles) -> None:
    """Questions (with options, no answers), a page break, then an answer key."""
    grouped = _group_by_type(questions)
    flow.append(Paragraph("Worksheet", styles["H2x"]))
    for qtype, items in grouped.items():
        flow.append(Paragraph(label_for(qtype), styles["H3x"]))
        for i, q in enumerate(items, 1):
            flow.append(Paragraph(f"<b>Q{i}.</b> {_inline_html(q.question)}", styles["Bodyx"]))
            opts = _options(q)
            if opts:
                flow.append(
                    ListFlowable(
                        [ListItem(Paragraph(_inline_html(o), styles["Bodyx"])) for o in opts],
                        bulletType="a",
                        leftIndent=18,
                    )
                )
            flow.append(Spacer(1, 10))

    flow.append(PageBreak())
    flow.append(Paragraph("Answer Key", styles["H2x"]))
    for qtype, items in grouped.items():
        flow.append(Paragraph(label_for(qtype), styles["H3x"]))
        for i, q in enumerate(items, 1):
            flow.append(
                Paragraph(f"<b>Q{i}.</b> {_inline_html(q.answer)}", styles["Bodyx"])
            )
            if q.explanation:
                flow.append(
                    Paragraph(f"<i>{_inline_html(q.explanation)}</i>", styles["Bodyx"])
                )
            flow.append(Spacer(1, 4))


def build_pdf(
    chapter, questions, include_notes: bool, include_questions: bool, worksheet: bool = False
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm,
    )
    styles = _pdf_styles()
    flow: list = [
        Paragraph(_inline_html(chapter.chapter_name), styles["H1x"]),
        Paragraph(
            f"Class: {chapter.class_name} &nbsp;|&nbsp; Subject: {chapter.subject}",
            styles["Metax"],
        ),
        Spacer(1, 10),
    ]

    if include_notes and chapter.notes.strip():
        flow.append(Paragraph("Notes", styles["H2x"]))
        flow += _markdown_to_flowables(clean_markdown_for_export(chapter.notes), styles)
        flow.append(Spacer(1, 12))

    if include_notes:
        sources = _sources(chapter)
        if sources:
            flow.append(Paragraph("References", styles["H2x"]))
            items = []
            for s in sources:
                title = _inline_html(s.get("title") or s["url"])
                href = s["url"].replace("&", "&amp;")
                items.append(
                    ListItem(
                        Paragraph(f'{title} — <a href="{href}">{href}</a>', styles["Bodyx"])
                    )
                )
            flow.append(ListFlowable(items, bulletType="1", leftIndent=18))
            flow.append(Spacer(1, 12))

    if include_questions and questions and worksheet:
        _pdf_worksheet(flow, questions, styles)
    elif include_questions and questions:
        flow.append(Paragraph("Question Bank", styles["H2x"]))
        for qtype, items in _group_by_type(questions).items():
            flow.append(Paragraph(label_for(qtype), styles["H3x"]))
            for i, q in enumerate(items, 1):
                flow.append(Paragraph(f"<b>Q{i}.</b> {_inline_html(q.question)}", styles["Bodyx"]))
                opts = _options(q)
                if opts:
                    flow.append(
                        ListFlowable(
                            [ListItem(Paragraph(_inline_html(o), styles["Bodyx"])) for o in opts],
                            bulletType="a",
                            leftIndent=18,
                        )
                    )
                flow.append(Paragraph(f"<b>Answer:</b> {_inline_html(q.answer)}", styles["Bodyx"]))
                if q.explanation:
                    flow.append(
                        Paragraph(
                            f"<i>Explanation: {_inline_html(q.explanation)}</i>", styles["Bodyx"]
                        )
                    )
                flow.append(Spacer(1, 6))

    if not flow[3:]:
        flow.append(Paragraph("Nothing to export yet.", styles["Bodyx"]))

    doc.build(flow)
    return buf.getvalue()


# ------------------------------- DOCX ------------------------------- #

def _docx_markdown(doc: Document, md: str) -> None:
    for raw in md.splitlines():
        stripped = raw.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            doc.add_heading(strip_md(stripped[4:]), level=3)
        elif stripped.startswith("## "):
            doc.add_heading(strip_md(stripped[3:]), level=2)
        elif stripped.startswith("# "):
            doc.add_heading(strip_md(stripped[2:]), level=1)
        elif stripped.startswith(("- ", "* ", "+ ")):
            doc.add_paragraph(strip_md(stripped[2:]), style="List Bullet")
        elif stripped.startswith("> "):
            run = doc.add_paragraph().add_run(strip_md(stripped[2:]))
            run.italic = True
        else:
            _add_rich_paragraph(doc, stripped)


def _add_rich_paragraph(doc: Document, text: str) -> None:
    """Add a paragraph rendering **bold** segments."""
    p = doc.add_paragraph()
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        p.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def _docx_worksheet(doc: Document, questions) -> None:
    grouped = _group_by_type(questions)
    doc.add_heading("Worksheet", level=1)
    for qtype, items in grouped.items():
        doc.add_heading(label_for(qtype), level=2)
        for i, q in enumerate(items, 1):
            _add_rich_paragraph(doc, f"**Q{i}.** {q.question}")
            for o in _options(q):
                doc.add_paragraph(strip_md(o), style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Answer Key", level=1)
    for qtype, items in grouped.items():
        doc.add_heading(label_for(qtype), level=2)
        for i, q in enumerate(items, 1):
            _add_rich_paragraph(doc, f"**Q{i}.** {q.answer}")
            if q.explanation:
                run = doc.add_paragraph().add_run(strip_md(q.explanation))
                run.italic = True


def build_docx(
    chapter, questions, include_notes: bool, include_questions: bool, worksheet: bool = False
) -> bytes:
    doc = Document()
    doc.add_heading(chapter.chapter_name, level=0)
    meta = doc.add_paragraph()
    run = meta.add_run(f"Class: {chapter.class_name}  |  Subject: {chapter.subject}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    if include_notes and chapter.notes.strip():
        doc.add_heading("Notes", level=1)
        _docx_markdown(doc, clean_markdown_for_export(chapter.notes))

    if include_notes:
        sources = _sources(chapter)
        if sources:
            doc.add_heading("References", level=1)
            for s in sources:
                title = strip_md(s.get("title") or s["url"])
                doc.add_paragraph(f"{title} — {s['url']}", style="List Number")

    if include_questions and questions and worksheet:
        _docx_worksheet(doc, questions)
    elif include_questions and questions:
        doc.add_heading("Question Bank", level=1)
        for qtype, items in _group_by_type(questions).items():
            doc.add_heading(label_for(qtype), level=2)
            for i, q in enumerate(items, 1):
                _add_rich_paragraph(doc, f"**Q{i}.** {q.question}")
                for o in _options(q):
                    doc.add_paragraph(strip_md(o), style="List Bullet")
                _add_rich_paragraph(doc, f"**Answer:** {q.answer}")
                if q.explanation:
                    run = doc.add_paragraph().add_run(f"Explanation: {strip_md(q.explanation)}")
                    run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
